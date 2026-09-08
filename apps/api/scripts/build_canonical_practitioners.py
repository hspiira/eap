"""Turn the reference workbook's canonical counselor list into an import file.

The activity log names counselors the practitioners workbook never listed, so
most of its sessions cannot be attributed. The `Counselors` sheet carries a
curated canonical list: one spelling per person, chosen by somebody who knows
them. This writes the names that are not yet practitioners into a workbook the
existing practitioner import reads, so they are created through the same
staged, reviewed, audited path as every other practitioner rather than by a
second route invented for them.

Names already held as practitioners are left out, so applying the batch cannot
produce a second record for one person. Nothing else is filled in: the list
supplies a name and nothing else, and a profession or an organisation that is
not in the source is not going to be guessed here.

    uv run --with openpyxl python scripts/build_canonical_practitioners.py \
        --workbook <path> --tenant-id <id>
"""

from __future__ import annotations

import argparse
import asyncio
import os
import sys
from pathlib import Path

SCRIPTS_DIR = Path(__file__).resolve().parent
if str(SCRIPTS_DIR) not in sys.path:
    sys.path.insert(0, str(SCRIPTS_DIR))

API_ROOT = SCRIPTS_DIR.parent
if str(API_ROOT) not in sys.path:
    sys.path.insert(0, str(API_ROOT))

from import_support import send  # noqa: E402

OUT = API_ROOT / "data" / "taxonomy" / "canonical_practitioners.xlsx"
#: The sheet name and headers the practitioner import already reads. It reads
#: both of the workbook's sheets and refuses a file missing either, so the
#: partner sheet is written out with its headers and no rows.
SHEET = "EAP Consultants - General"
HEADERS = ("NAME", "COMPANY", "Speciality", "EMAIL", "Contract")
PARTNER_SHEET = "Minet EAP Partner list"
PARTNER_HEADERS = ("NAME (First/Surname)", "INDIVIDUAL/COMPANY NAME", "PROFESSION", "CONTACT EMAIL")
#: Column D of the Counselors sheet: the curated one-per-person list.
CANONICAL_COLUMN = 3


def _canonical(workbook: Path, sheet: str) -> list[str]:
    import openpyxl

    rows = openpyxl.load_workbook(workbook, read_only=True)[sheet].iter_rows(
        min_row=2, values_only=True
    )
    seen: dict[str, str] = {}
    for row in rows:
        if len(row) > CANONICAL_COLUMN and row[CANONICAL_COLUMN]:
            name = str(row[CANONICAL_COLUMN]).strip()
            seen.setdefault(name.casefold(), name)
    return sorted(seen.values())


async def _existing(tenant_id: str) -> set[str]:
    """Normalised names already held as practitioners in the environment."""
    from httpx import ASGITransport, AsyncClient

    from app.core.security import TokenData, get_current_user, get_current_user_optional
    from app.domain.services.provider_alias_normalisation import normalise_practitioner_name
    from app.main import app

    actor = TokenData(user_id=os.environ["IMPORT_ACTOR_ID"], tenant_id=tenant_id, role="Admin")
    app.dependency_overrides[get_current_user] = lambda: actor
    app.dependency_overrides[get_current_user_optional] = lambda: actor

    names: set[str] = set()
    async with AsyncClient(
        transport=ASGITransport(app=app), base_url="http://import", timeout=300
    ) as http:
        page = 1
        while True:
            response = await send(
                http, "GET", f"/providers?tenant_id={tenant_id}&page={page}&limit=100"
            )
            response.raise_for_status()
            body = response.json()
            names.update(normalise_practitioner_name(p["display_name"]) for p in body["items"])
            if not body.get("has_more"):
                return names
            page += 1


def _write(names: list[str]) -> None:
    import openpyxl

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = SHEET
    sheet.append(list(HEADERS))
    for name in names:
        sheet.append([name, None, None, None, None])
    partner = workbook.create_sheet(PARTNER_SHEET)
    partner.append([])
    partner.append(list(PARTNER_HEADERS))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)


async def run(workbook: Path, sheet: str, tenant_id: str) -> int:
    from app.domain.services.provider_alias_normalisation import normalise_practitioner_name

    canonical = _canonical(workbook, sheet)
    held = await _existing(tenant_id)
    missing = [name for name in canonical if normalise_practitioner_name(name) not in held]
    _write(missing)

    print(f"{OUT.name:28} {len(missing)} of {len(canonical)} canonical names are missing")
    for name in canonical:
        if normalise_practitioner_name(name) in held:
            print(f"  already a practitioner: {name}")
    partial = [name for name in missing if len(normalise_practitioner_name(name).split()) < 2]
    if partial:
        print(
            f"  {len(partial)} of these are one word and cannot tell two people apart: "
            + ", ".join(partial)
        )
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--workbook", required=True, type=Path)
    parser.add_argument("--sheet", default="Counselors")
    parser.add_argument("--tenant-id", required=True)
    args = parser.parse_args()
    if not os.environ.get("IMPORT_ACTOR_ID"):
        sys.exit("IMPORT_ACTOR_ID must name the admin the import is recorded against.")
    os.environ.setdefault("PYTHONWARNINGS", "ignore")
    return asyncio.run(run(args.workbook, args.sheet, args.tenant_id))


if __name__ == "__main__":
    raise SystemExit(main())
